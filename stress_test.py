#!/usr/bin/env python3
"""
14-day simulated deployment stress test for TTTTMTTRT.

Runs 5 mock job postings/day x 14 days = 70 total tailoring runs against
master_resume.docx through the real Gemini tailoring engine, then validates
domain relevance, exclusion of incompatible content, and structural integrity
of the output before it ever gets published.
"""

import random
import time

import colorama
from colorama import Fore, Style
from docx import Document

import main as app

colorama.init(autoreset=True)

DAYS = 14
JOBS_PER_DAY = 5
THROTTLE_SECONDS = 1.5

COMPANIES = [
    "Northgate Legal Partners", "Vantage Law Group", "Sterling & Cross LLP",
    "Bluepeak Administrative Services", "Harborview Operations Co.",
    "Crestline Business Solutions", "Nimbus Software Inc.", "Quantum Byte Labs",
    "Redwood Dev Studio", "Marble Bay Customer Care", "Lighthouse Support Co.",
    "Fernwood Client Services", "Dominion Capital Partners", "Ashford Finance Group",
    "Ledgerline Corporate Advisors",
]

DOMAINS = {
    "legal": {
        "keywords": ["legal research", "case law", "litigation", "statutory analysis", "policy memo"],
        "incompatible": ["Python", "React", "API", "database schema", "unit testing"],
        "title": "Legal Research Assistant",
    },
    "administrative": {
        "keywords": ["workflow coordination", "client service", "scheduling", "process automation", "office operations"],
        "incompatible": ["case law", "litigation", "statutory analysis", "unit testing", "React"],
        "title": "Operations Coordinator",
    },
    "software engineering": {
        "keywords": ["Python", "REST API", "unit testing", "database schema", "version control"],
        "incompatible": ["case law", "litigation", "statutory analysis"],
        "title": "Software Developer Co-op",
    },
    "customer service": {
        "keywords": ["client satisfaction", "ticket resolution", "communication skills", "CRM software", "escalation handling"],
        "incompatible": ["case law", "database schema", "unit testing"],
        "title": "Customer Service Representative",
    },
    "corporate finance": {
        "keywords": ["financial modeling", "budget forecasting", "variance analysis", "reporting accuracy", "Excel"],
        "incompatible": ["case law", "React", "unit testing"],
        "title": "Corporate Finance Analyst",
    },
}


def generate_job_descriptions():
    jobs = []
    domain_names = list(DOMAINS.keys())
    company_pool = COMPANIES[:]
    idx = 0
    for day in range(1, DAYS + 1):
        for domain in domain_names:
            spec = DOMAINS[domain]
            company = company_pool[idx % len(company_pool)]
            idx += 1
            description = (
                f"{company} is hiring a {spec['title']} for a 4-month co-op term.\n"
                f"Key responsibilities and requirements include: {', '.join(spec['keywords'])}.\n"
                f"Ideal candidates demonstrate strength in {spec['keywords'][0]} and {spec['keywords'][-1]}."
            )
            jobs.append({
                "day": day,
                "domain": domain,
                "company": company,
                "description": description,
                "keywords": spec["keywords"],
                "incompatible": spec["incompatible"],
            })
    return jobs


def extract_all_text(doc):
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(parts)


def validate(resume_text, job):
    reasons = []

    hits = sum(1 for kw in job["keywords"] if kw.lower() in resume_text.lower())
    domain_relevance_pass = hits >= max(1, len(job["keywords"]) // 2)
    if not domain_relevance_pass:
        reasons.append(f"domain relevance low ({hits}/{len(job['keywords'])} keywords found)")

    leaked = [kw for kw in job["incompatible"] if kw.lower() in resume_text.lower()]
    exclusion_pass = len(leaked) == 0
    if not exclusion_pass:
        reasons.append(f"incompatible content leaked: {leaked}")

    bad_tokens = [app.REMOVE_TOKEN, "<w:", "</w:", "{{", "}}", "TRSM", "Ryerson"]
    found_bad = [t for t in bad_tokens if t in resume_text]
    structural_pass = len(found_bad) == 0
    if not structural_pass:
        reasons.append(f"structural integrity failed: {found_bad}")

    passed = domain_relevance_pass and exclusion_pass and structural_pass
    return passed, reasons


def run_stress_test():
    print(f"{Fore.CYAN}RUNNING 14-DAY STRESS TEST (70 JOB POSTINGS)...{Style.RESET_ALL}")

    if not (app.os.path.exists(app.MASTER_RESUME)):
        print(f"{Fore.RED}Missing {app.MASTER_RESUME} in current folder. Cannot run stress test.{Style.RESET_ALL}")
        return

    client = app.get_client()
    jobs = generate_job_descriptions()

    results = []
    for i, job in enumerate(jobs, start=1):
        try:
            resume_doc = Document(app.MASTER_RESUME)
            resume_map = app.build_paragraph_map(resume_doc)
            result = app.request_edits(client, app.RESUME_RULES, job["description"], resume_map)
            app.apply_edits(resume_doc, result["edits"])
            resume_text = extract_all_text(resume_doc)
            passed, reasons = validate(resume_text, job)
        except Exception as exc:
            passed, reasons = False, [f"exception: {exc}"]

        status = f"{Fore.GREEN}PASS{Style.RESET_ALL}" if passed else f"{Fore.RED}FAIL{Style.RESET_ALL}"
        print(f"Day {job['day']:>2} | {job['domain']:<20} | {job['company']:<28} | {status}"
              + (f" -- {'; '.join(reasons)}" if reasons else ""))

        results.append({**job, "passed": passed, "reasons": reasons})
        time.sleep(THROTTLE_SECONDS)

    total = len(results)
    passed_count = sum(1 for r in results if r["passed"])

    print("\n" + "=" * 60)
    print(f"{Fore.YELLOW}STRESS TEST SUMMARY{Style.RESET_ALL}")
    print("=" * 60)
    for domain in DOMAINS:
        domain_results = [r for r in results if r["domain"] == domain]
        domain_passed = sum(1 for r in domain_results if r["passed"])
        print(f"{domain:<20}: {domain_passed}/{len(domain_results)} PASSED")
    print("-" * 60)
    color = Fore.GREEN if passed_count == total else Fore.RED
    print(f"{color}TOTAL: {passed_count}/{total} PASSED{Style.RESET_ALL}")
    print("=" * 60)


if __name__ == "__main__":
    run_stress_test()
